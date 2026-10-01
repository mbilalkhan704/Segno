"""Pure helpers (no Tk): Meraki-compatible filenames, CSV/Excel reading,
email validation, certificate-folder matching, row analysis, atomic JSON
writes, and Windows DPAPI encryption for the stored app password."""

import base64
import csv
import difflib
import json
import os
import re
import sys
import ctypes
from dataclasses import dataclass, field
from typing import Optional

from cm_constants import (
    CERT_EXTENSIONS, ID_MATCH_MIN_NAME_SIMILARITY,
)


# ---------------------------------------------------------------------------
# Filenames - verbatim from Meraki's mck_utils so matching stays consistent
# ---------------------------------------------------------------------------
def sanitize_filename(name: str) -> str:
    name = name.strip()
    name = re.sub(r'[\\/*?:"<>|]', "", name)
    name = re.sub(r"\s+", " ", name)
    return name or "unnamed"


def name_to_filename(raw_name: str) -> str:
    """'Muhammad bilal khan' -> 'Muhammad_Bilal_Khan'."""
    cleaned = sanitize_filename(raw_name)
    words = [w.capitalize() for w in cleaned.split(" ") if w]
    return "_".join(words) or "Unnamed"


def cert_stem(name: str, cert_id: str) -> str:
    """The filename stem Meraki writes: name, plus _ID only when an ID exists."""
    base = name_to_filename(name)
    cid = (cert_id or "").strip()
    return f"{base}_{sanitize_filename(cid)}" if cid else base


def make_row_key(name: str, cert_id: str) -> str:
    """Identity used by the sent log: the Certificate ID when present."""
    cid = (cert_id or "").strip()
    return cid if cid else "name:" + name_to_filename(name)


# ---------------------------------------------------------------------------
# Tabular file reading (Meraki's reader + fixes: None->"", closed workbook,
# duplicate/blank headers, trailing empty Excel columns, encoding fallback)
# ---------------------------------------------------------------------------
def _unique_headers(raw_headers):
    seen, out = set(), []
    for i, h in enumerate(raw_headers):
        s = str(h).strip() if h is not None else ""
        if not s:
            s = f"Column {i + 1}"
        if s in seen:
            n = 2
            while f"{s} ({n})" in seen:
                n += 1
            s = f"{s} ({n})"
        seen.add(s)
        out.append(s)
    return out


def _cell_to_str(val) -> str:
    if val is None:
        return ""
    if isinstance(val, float) and val.is_integer():
        val = int(val)
    return str(val)


def read_tabular_file(path: str):
    """Return (headers, rows) for a .csv/.xlsx/.xlsm; rows are dicts keyed by
    header with every value a string. Raises on any problem."""
    ext = os.path.splitext(path)[1].lower()

    if ext == ".csv":
        raw_rows = None
        for enc in ("utf-8-sig", "cp1252"):
            try:
                with open(path, newline="", encoding=enc) as f:
                    raw_rows = list(csv.reader(f))
                break
            except UnicodeDecodeError:
                continue
        if raw_rows is None:
            raise ValueError("Couldn't decode the CSV file (try saving it as UTF-8).")
        raw_rows = [r for r in raw_rows if r]          # skip fully blank lines
        if not raw_rows:
            return [], []
        headers = _unique_headers(raw_rows[0])
        rows = []
        for raw in raw_rows[1:]:
            rows.append({h: (raw[i] if i < len(raw) else "") for i, h in enumerate(headers)})
        return headers, rows

    if ext in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook  # type: ignore
        except ImportError:
            raise RuntimeError("Reading Excel files needs the 'openpyxl' package.\n"
                               "Install it with:  pip install openpyxl")
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
            try:
                header_row = list(next(rows_iter))
            except StopIteration:
                return [], []
            last = -1
            for i, h in enumerate(header_row):
                if h is not None and str(h).strip():
                    last = i
            if last < 0:
                return [], []
            header_row = header_row[:last + 1]
            headers = _unique_headers(header_row)
            rows = []
            for raw_row in rows_iter:
                if raw_row is None or all(v is None for v in raw_row):
                    continue
                rows.append({h: _cell_to_str(raw_row[i]) if i < len(raw_row) else ""
                             for i, h in enumerate(headers)})
            return headers, rows
        finally:
            wb.close()

    raise ValueError(f"Unsupported file type: {ext or '(none)'}. Use .csv or .xlsx.")


def guess_column(headers, candidates, contains=()):
    """Auto-pick a column by header name (exact match first, then 'contains')."""
    lowered = {h: h.strip().lower() for h in headers}
    for cand in candidates:
        for h, low in lowered.items():
            if low == cand:
                return h
    for frag in contains:
        for h, low in lowered.items():
            if frag in low:
                return h
    return ""


# ---------------------------------------------------------------------------
# Email validation (pragmatic; single address only)
# ---------------------------------------------------------------------------
_EMAIL_RE = re.compile(
    r"^[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@"
    r"[A-Za-z0-9](?:[A-Za-z0-9-]{0,61}[A-Za-z0-9])?"
    r"(?:\.[A-Za-z0-9](?:[A-Za-z0-9-]{0,61}[A-Za-z0-9])?)+$")


def clean_email(raw) -> str:
    s = (raw or "").replace("\u00a0", " ").replace("\u200b", "").strip()
    if s.lower().startswith("mailto:"):
        s = s[7:]
    return s.strip(" <>\t\r\n")


def validate_email(raw):
    """Return (ok, cleaned, reason). reason is '' when ok."""
    s = clean_email(raw)
    if not s:
        return False, "", "empty"
    if re.search(r"[;,\s]", s):
        return False, s, "multiple or malformed"
    local, _, domain = s.rpartition("@")
    if len(s) > 254 or len(local) > 64 or not _EMAIL_RE.match(s):
        return False, s, "invalid"
    return True, s, ""


# ---------------------------------------------------------------------------
# Certificate folder index + matching
# ---------------------------------------------------------------------------
_EXT_RANK = {ext: i for i, ext in enumerate(CERT_EXTENSIONS)}


def _norm_id(s: str) -> str:
    s = s.strip().lower()
    return (s.lstrip("0") or "0") if s.isdigit() else s


class CertIndex:
    """One scan of the certificates folder; lookups are O(1) per row."""

    def __init__(self, folder):
        self.folder = folder
        self.by_stem = {}     # lower stem -> {ext: path}
        self.by_id = {}       # normalized id -> {lower stem: {ext: path}}
        self.count = 0
        self.error = None
        try:
            entries = list(os.scandir(folder))
        except OSError as exc:
            self.error = str(exc)
            return
        for e in entries:
            try:
                if not e.is_file():
                    continue
            except OSError:
                continue
            stem, ext = os.path.splitext(e.name)
            ext = ext.lower()
            if ext not in _EXT_RANK:
                continue
            self.count += 1
            low = stem.lower()
            self.by_stem.setdefault(low, {})[ext] = e.path
            if "_" in stem:
                self.by_id.setdefault(_norm_id(stem.rsplit("_", 1)[1]), {}) \
                    .setdefault(low, {})[ext] = e.path

    @staticmethod
    def _pick(ext_map):
        return ext_map[min(ext_map, key=_EXT_RANK.__getitem__)]

    def find(self, name, cert_id):
        """Return (path, how) where how is 'exact', 'id', or None."""
        name = (name or "").strip()
        cid = (cert_id or "").strip()
        if not name:
            return None, None
        expected = cert_stem(name, cid).lower()
        exact = self.by_stem.get(expected)
        if exact:
            return self._pick(exact), "exact"
        if cid:
            stems = self.by_id.get(_norm_id(sanitize_filename(cid)), {})
            if len(stems) == 1:
                stem, ext_map = next(iter(stems.items()))
                name_part = stem.rsplit("_", 1)[0]
                want = name_to_filename(name).lower()
                if difflib.SequenceMatcher(None, name_part, want).ratio() >= ID_MATCH_MIN_NAME_SIMILARITY:
                    return self._pick(ext_map), "id"
        return None, None


@dataclass
class RowInfo:
    index: int                      # 1-based, as shown in the grid
    name: str = ""
    email: str = ""
    email_ok: bool = False
    email_reason: str = ""
    cert_id: str = ""
    cert_path: Optional[str] = None
    match_how: Optional[str] = None
    key: str = ""
    dup_email: bool = False
    cert_reason: str = ""           # '', 'no name', 'no certificate'

    @property
    def sendable(self) -> bool:
        return self.email_ok and self.cert_path is not None


def analyze_rows(rows, name_col, email_col, cert_index, cert_id_col):
    """Build one RowInfo per row: email validity + certificate availability."""
    infos = []
    for i, row in enumerate(rows, start=1):
        name = (row.get(name_col) or "").strip() if name_col else ""
        cid = (row.get(cert_id_col) or "").strip()
        ok, email, reason = validate_email(row.get(email_col) if email_col else "")
        info = RowInfo(index=i, name=name, email=email, email_ok=ok, email_reason=reason,
                       cert_id=cid, key=make_row_key(name, cid))
        if not name:
            info.cert_reason = "no name"
        elif cert_index is None or cert_index.error:
            info.cert_reason = "no certificate"
        else:
            path, how = cert_index.find(name, cid)
            if path:
                info.cert_path, info.match_how = path, how
            else:
                info.cert_reason = "no certificate"
        infos.append(info)
    counts = {}
    for info in infos:
        if info.email_ok:
            counts[info.email.lower()] = counts.get(info.email.lower(), 0) + 1
    for info in infos:
        info.dup_email = info.email_ok and counts[info.email.lower()] > 1
    return infos


# ---------------------------------------------------------------------------
# JSON + DPAPI
# ---------------------------------------------------------------------------
def atomic_write_json(path, data):
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=1)
    os.replace(tmp, path)


def _dpapi_call(data: bytes, encrypt: bool) -> bytes:
    from ctypes import wintypes

    class DATA_BLOB(ctypes.Structure):
        _fields_ = [("cbData", wintypes.DWORD), ("pbData", ctypes.POINTER(ctypes.c_char))]

    buf = ctypes.create_string_buffer(data, len(data))
    blob_in = DATA_BLOB(len(data), ctypes.cast(buf, ctypes.POINTER(ctypes.c_char)))
    blob_out = DATA_BLOB()
    fn = ctypes.windll.crypt32.CryptProtectData if encrypt else ctypes.windll.crypt32.CryptUnprotectData
    ok = fn(ctypes.byref(blob_in), None, None, None, None, 0, ctypes.byref(blob_out))
    if not ok:
        raise OSError("DPAPI call failed")
    try:
        return ctypes.string_at(blob_out.pbData, blob_out.cbData)
    finally:
        ctypes.windll.kernel32.LocalFree(ctypes.cast(blob_out.pbData, ctypes.c_void_p))


def protect_secret(text: str) -> str:
    """Encrypt for storage. Windows: DPAPI (tied to this Windows user).
    Elsewhere (dev/tests only): base64, clearly prefixed as NOT encrypted."""
    if not text:
        return ""
    raw = text.encode("utf-8")
    if sys.platform.startswith("win"):
        try:
            return "dpapi:" + base64.b64encode(_dpapi_call(raw, True)).decode("ascii")
        except Exception:
            pass
    return "plain:" + base64.b64encode(raw).decode("ascii")


def unprotect_secret(stored: str) -> str:
    if not stored:
        return ""
    try:
        kind, _, payload = stored.partition(":")
        raw = base64.b64decode(payload)
        if kind == "dpapi":
            raw = _dpapi_call(raw, False)
        elif kind != "plain":
            return ""
        return raw.decode("utf-8")
    except Exception:
        return ""
