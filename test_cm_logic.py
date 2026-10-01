"""Logic tests (no GUI). Run:  python -m unittest test_cm_logic -v"""
import os, tempfile, threading, unittest, smtplib, email

os.environ["CM_APPDATA_DIR"] = tempfile.mkdtemp()

from cm_utils import (name_to_filename, cert_stem, validate_email, CertIndex, analyze_rows,
                      read_tabular_file, protect_secret, unprotect_secret, guess_column)
from cm_sentlog import SentLog
from cm_richtext import (make_run, substitute_name, runs_to_html, runs_to_plain, normalize_runs,
                         find_unknown_placeholders, normalize_url)
from cm_mail import render_email, build_message, classify_exception, run_batch, validate_smtp_credentials
from cm_constants import CERT_ID_COLUMN


def write_bytes(path, data=b"%PDF-1.4 test"):
    with open(path, "wb") as f:
        f.write(data)


class FakeSMTP:
    instances = []
    fail_login = False
    send_behaviour = None     # callable(msg, n) -> may raise
    attempts = 0
    def __init__(self, host, port, timeout=None):
        self.sent = []; self.connected = True; FakeSMTP.instances.append(self)
    def ehlo(self): pass
    def starttls(self, context=None): pass
    def login(self, u, p):
        if FakeSMTP.fail_login:
            raise smtplib.SMTPAuthenticationError(535, b"5.7.8 Username and Password not accepted")
    def send_message(self, msg):
        n = FakeSMTP.attempts; FakeSMTP.attempts += 1
        if FakeSMTP.send_behaviour: FakeSMTP.send_behaviour(msg, n)
        self.sent.append(msg)
    def quit(self): self.connected = False
    def close(self): self.connected = False


class Files(unittest.TestCase):
    def test_stem_matches_meraki(self):
        self.assertEqual(name_to_filename("Muhammad bilal khan"), "Muhammad_Bilal_Khan")
        self.assertEqual(cert_stem("Muhammad bilal khan", "4K9P2Q"), "Muhammad_Bilal_Khan_4K9P2Q")
        self.assertEqual(cert_stem("Ali", ""), "Ali")
        self.assertEqual(name_to_filename('A/B:C  d'), "Abc_D")

    def test_matching_and_fallbacks(self):
        d = tempfile.mkdtemp()
        for f in ("Ali_Khan_AB12CD.png", "Ali_Khan_AB12CD.pdf", "Sara_Noor_004917.pdf",
                  "Usman_Raza_ZZ99ZZ.jpg", "notes.txt"):
            write_bytes(os.path.join(d, f), b"x")
        idx = CertIndex(d)
        p, how = idx.find("ali khan", "ab12cd")            # case-insensitive, pdf preferred
        self.assertTrue(p.endswith(".pdf") and how == "exact")
        p, how = idx.find("Sara Noor", "4917")             # lost leading zeros
        self.assertEqual(how, "id")
        p, how = idx.find("Usmaan Raza", "ZZ99ZZ")         # typo fixed after generation
        self.assertEqual(how, "id")
        p, how = idx.find("Totally Different", "ZZ99ZZ")   # same ID, other person -> refuse
        self.assertIsNone(p)
        self.assertEqual(idx.find("", "AB12CD"), (None, None))
        self.assertIsNotNone(CertIndex(os.path.join(d, "nope")).error)

    def test_analyze(self):
        d = tempfile.mkdtemp(); write_bytes(os.path.join(d, "Ali_Khan_A1.pdf"), b"x")
        rows = [{"Name": "Ali Khan", "Email": "ali@x.com", CERT_ID_COLUMN: "A1"},
                {"Name": "Bob", "Email": "bad", CERT_ID_COLUMN: "B2"},
                {"Name": "", "Email": "a@b.co", CERT_ID_COLUMN: ""},
                {"Name": "Ali Khan", "Email": "ALI@x.com", CERT_ID_COLUMN: "A1"}]
        infos = analyze_rows(rows, "Name", "Email", CertIndex(d), CERT_ID_COLUMN)
        self.assertTrue(infos[0].sendable and infos[0].dup_email)
        self.assertFalse(infos[1].sendable)
        self.assertEqual(infos[2].cert_reason, "no name")

    def test_emails(self):
        self.assertTrue(validate_email(" Ali@Example.com\u00a0")[0])
        self.assertTrue(validate_email("mailto:a.b+c@uok.edu.pk")[0])
        for bad in ("", "x", "a@b", "a@b.com; c@d.com", "a b@c.com", "@x.com"):
            self.assertFalse(validate_email(bad)[0], bad)

    def test_read_csv_and_excel(self):
        d = tempfile.mkdtemp(); p = os.path.join(d, "a.csv")
        write_bytes(p, "Name,Name,Email\nAli,x\n\nSara,y,s@x.com\n".encode("utf-8-sig"))
        h, r = read_tabular_file(p)
        self.assertEqual(h, ["Name", "Name (2)", "Email"]); self.assertEqual(len(r), 2)
        self.assertEqual(r[0]["Email"], "")                # short row -> "" not None
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active; ws.append(["Name", "Email", None, None]); ws.append(["A", "a@b.co", None, None]); ws.append([None] * 4)
        x = os.path.join(d, "a.xlsx"); wb.save(x)
        h, r = read_tabular_file(x); self.assertEqual(h, ["Name", "Email"]); self.assertEqual(len(r), 1)
        self.assertEqual(guess_column(["Full Name", "E-mail"], ("email", "e-mail")), "E-mail")

    def test_secret_roundtrip(self):
        self.assertEqual(unprotect_secret(protect_secret("abcd efgh")), "abcd efgh")
        self.assertEqual(unprotect_secret("garbage"), "")

    def test_sentlog(self):
        log = SentLog(os.path.join(tempfile.mkdtemp(), "s.json"))
        log.record("K1", "a@x.com", "Ali", "f.pdf")
        self.assertTrue(log.is_sent("K1", "A@X.com")); self.assertFalse(log.is_sent("K1", "new@x.com"))
        self.assertEqual(SentLog(log.path).count_since(24), 1)


class RichText(unittest.TestCase):
    def test_substitute_across_runs_and_escape(self):
        runs = [make_run("Dear {na"), make_run("me}", b=True), make_run(", hi")]
        out = substitute_name(runs, "A<b>&Co")
        self.assertEqual("".join(r["text"] for r in out), "Dear A<b>&Co, hi")
        self.assertIn("A&lt;b&gt;&amp;Co", runs_to_html(out))

    def test_html_and_plain(self):
        runs = [make_run("Hi ", b=True), make_run("site", href="https://x.org", u=True), make_run("big", size=20)]
        h = runs_to_html(runs)
        self.assertIn("<b>Hi </b>", h); self.assertIn('href="https://x.org"', h); self.assertIn("font-size:20px", h)
        self.assertEqual(runs_to_plain(runs), "Hi site (https://x.org)big")

    def test_fonts(self):
        r = [make_run("Hi ", font="Georgia"), make_run("there", font="Georgia", size=20, b=True),
             make_run(" x", font="Papyrus"), make_run(" y", font="Arial")]
        out = normalize_runs(r)
        self.assertEqual([x["font"] for x in out], ["Georgia", "Georgia", None])        # unknown/default -> None
        self.assertEqual(out[2]["text"], " x y")                                       # merged: both default now
        h = runs_to_html(r)
        self.assertIn("font-family:Georgia, 'Times New Roman', serif;", h)
        self.assertIn("font-family:Georgia, 'Times New Roman', serif;font-size:20px;", h)
        self.assertNotIn("Papyrus", h); self.assertNotIn("Arial, Helvetica", h)
        sub = substitute_name([make_run("Dear {name}", font="Verdana")], "Ali")
        self.assertEqual((sub[0]["text"], sub[0]["font"]), ("Dear Ali", "Verdana"))
        self.assertEqual(normalize_runs([{"text": "old saved run", "b": False}])[0]["font"], None)  # old data

    def test_misc(self):
        self.assertEqual(find_unknown_placeholders("Hi {name} {Nmae}"), ["{Nmae}"])
        self.assertIsNone(normalize_url("javascript:alert(1)"))
        self.assertEqual(normalize_url("uok.edu.pk"), "https://uok.edu.pk")
        self.assertEqual(normalize_url("a@b.com"), "mailto:a@b.com")
        self.assertEqual(len(normalize_runs([make_run("a"), make_run("b"), make_run("", b=True)])), 1)


class Mail(unittest.TestCase):
    cfg = {"email": "me@gmail.com", "password": "pw", "sender_name": "ORIC \u00c9"}
    def setUp(self):
        FakeSMTP.instances = []; FakeSMTP.fail_login = False; FakeSMTP.send_behaviour = None; FakeSMTP.attempts = 0
        d = tempfile.mkdtemp(); self.files = []
        for n in range(4):
            p = os.path.join(d, f"P{n}.pdf"); write_bytes(p); self.files.append(p)
        self.jobs = [{"row": n + 1, "name": f"P{n}", "email": f"p{n}@x.com", "cert_path": self.files[n]} for n in range(4)]
        self.render = lambda name: render_email("Your certificate, {name}", [make_run("Dear {name},")], [make_run("ORIC")], name)
        self.events = []

    def run_it(self, cancel=None, delay=0):
        run_batch(self.jobs, self.cfg, self.render, self.events.append, cancel or threading.Event(), delay, FakeSMTP)
        return [e[0] for e in self.events]

    def test_message_structure(self):
        subj, plain, html = self.render("Ali")
        m = build_message("me@gmail.com", "ORIC \u00c9", "a@x.com", subj, plain, html, self.files[0])
        parsed = email.message_from_bytes(bytes(m))
        self.assertEqual(parsed.get_content_type(), "multipart/mixed")
        kinds = [p.get_content_type() for p in parsed.walk()]
        self.assertIn("text/plain", kinds); self.assertIn("text/html", kinds); self.assertIn("application/pdf", kinds)
        self.assertEqual(parsed["Subject"], "Your certificate, Ali")

    def test_happy_path_single_connection(self):
        kinds = self.run_it()
        self.assertEqual(kinds.count("sent"), 4); self.assertEqual(len(FakeSMTP.instances), 1)
        self.assertEqual(kinds[-1], "done")

    def test_quota_aborts_and_marks_rest(self):
        def boom(msg, n):
            if n == 1: raise smtplib.SMTPDataError(550, b"5.4.5 Daily user sending limit exceeded")
        FakeSMTP.send_behaviour = boom
        kinds = self.run_it()
        self.assertEqual(kinds.count("sent"), 1); self.assertEqual(kinds.count("failed"), 1)
        self.assertEqual(kinds.count("skipped"), 2); self.assertIn("fatal", kinds)
        self.assertEqual([e for e in self.events if e[0] == "failed"][0][2], "quota")

    def test_bad_recipient_continues(self):
        def boom(msg, n):
            if n == 0: raise smtplib.SMTPRecipientsRefused({"p0@x.com": (550, b"no such user")})
        FakeSMTP.send_behaviour = boom
        kinds = self.run_it()
        self.assertEqual(kinds.count("sent"), 3); self.assertNotIn("fatal", kinds)

    def test_auth_failure_is_fatal_before_sending(self):
        FakeSMTP.fail_login = True
        kinds = self.run_it()
        self.assertEqual(kinds.count("skipped"), 4); self.assertEqual(kinds.count("sent"), 0)
        self.assertEqual(validate_smtp_credentials("a", "b", smtp_factory=FakeSMTP)[0], "invalid")

    def test_cancel(self):
        ev = threading.Event(); ev.set()
        self.assertEqual(self.run_it(cancel=ev).count("skipped"), 4)

    def test_oversize_and_missing_file(self):
        self.jobs[0]["cert_path"] = "/nonexistent.pdf"
        kinds = self.run_it(); self.assertEqual(kinds.count("failed"), 1); self.assertEqual(kinds.count("sent"), 3)

    def test_classify(self):
        self.assertEqual(classify_exception(smtplib.SMTPServerDisconnected("x"))[0], "connection")
        self.assertEqual(classify_exception(smtplib.SMTPSenderRefused(552, b"5.3.4 size", "a"))[0], "size")
        self.assertEqual(classify_exception(smtplib.SMTPDataError(421, b"4.7.0 slow down"))[0], "temporary")


if __name__ == "__main__":
    unittest.main()
